"""One-time derivation of locked-resume-template.docx from the UC layout sample.

Layout only: replace identity/content with join-contract placeholders and strip
Resume-Now custom properties. Does not invent jobs, skills, titles, dates, or metrics.
"""

from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import re

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

REPO_ROOT = Path(__file__).resolve().parents[2]
MAX_PRINTED_ROLES = 10
JOB_SLOT_TOKEN_RE = re.compile(r"\{\{JOB_HISTORY_(\d+)_")
RIGHT_COLUMN_TOPBORDER_STYLE = "divdocumentdivright-boxdisplaytabledisplaycelltopborder"
SAMPLE = REPO_ROOT / "scripts/resume-docx/tests/fixtures/layout-sample.docx"
OUTPUT = REPO_ROOT / ".ai/guardrails/locked-resume-template.docx"

ROLE_HEADER_MARKERS = [
    "Team Lead / Acting Enterprise Architect",
    "Team Lead / Acting Solution Architect",
    "Senior Analyst Enrollment Systems",
    "Associate Application Analyst",
    "Information Technology Systems Analyst",
]
ROLE_LOCATIONS = ["New York, New York"] * 5
ROLE_BULLETS = [
    [
        "Databricks Data Architecture: Engineered a Databricks-to-AWS Redshift proof-of-concept ingestion pipeline with governed field mappings, schema policies, and enterprise architecture controls.",
        "Lineage & Impact Analysis: Authored application catalogs, integration matrices, and mapping documents used by cross-functional teams to evaluate downstream system effects.",
        "Pipeline Recovery: Reverse-engineered Source System-to-Redshift-to-DBT-to-Looker reporting flows, traced fields to source schemas, resolved data gaps, and delivered stalled dashboards at go-live.",
    ],
    [
        "Migration Execution: Led the Salesforce Education Cloud implementation and completed end-to-end data migrations for student, primary-contact, and discipline-incident records.",
        "Transformation Design: Built data warehouse-to-Salesforce pipelines, governed cross-system model changes, and established interim templates during long-term ETL development.",
        "Engineering Enablement: Transitioned third-party ETL operations to an internal ownership model, defined technical standards, and mentored team members responsible for continued support.",
    ],
    [
        "Data Reconciliation: Engineered a VBA extraction and transformation utility that enabled sibling-record validation before SIS enrollment and reduced orphaned records by 90%.",
        "Scalable Integration: Directed technical patterns for the cross-functional student onboarding group and standardized data movement across connected cloud infrastructure during peak cycles",
        "Production Reliability: Delivered tier-3 support across the enterprise application landscape, lowering annual data-pipeline failure rates from 25% to under 5%.",
    ],
    [
        "Data Model Assessment: Evaluated enterprise warehouse and custom-development initiatives, mapped complex information flows, and partnered with engineers to improve reporting structures.",
        "Source-System Expertise: Served as SME for the Student Information System, warehouse models, and downstream BI reporting while translating constraints for business stakeholders.",
        "Quality & Ingestion Standards: Authored duplicate-record cleanup practices, reusable load templates, and documented procedures for external-system ingestion into the SIS.",
    ],
    [
        "SAP S/4HANA Consolidation: Served as retail systems technical lead for a global, multi-brand ERP consolidation, defining data-migration strategies and integration patterns for one enterprise platform.",
        "Cross-Platform Data Design: Supported synchronization of core structures and connected supply-chain repositories, improving consistency and operational reporting across global retail brands.",
        "Migration Documentation: Developed logical process flows and technical design specifications for global engineering teams, supporting cross-functional alignment and testing.",
    ],
]
SAMPLE_SKILLS = [
    "SAP S/4HANA",
    "Databricks",
    "SQL",
    "Data Migration",
    "Extract, Transform, Load (ETL)",
    "Data Architecture",
    "Systems Integration",
    "Enterprise Resource Planning (ERP)",
    "Solution Architecture",
    "Reverse Engineering",
    "Requirements Analysis",
    "Architecture Governance",
    "Technical Leadership",
    "Enterprise Architecture",
    "Software Architecture",
    "System Architecture",
    "Application Rationalization",
    "Technical Roadmaps",
    "Cloud Strategy",
    "Amazon Web Services (AWS)",
    "Strategic IT Alignment",
    "TOGAF",
    "Continuous Process Improvement",
    "Process Automation",
    "Project Management",
    "IT Operations Management",
    "Software Development Lifecycle (SDLC)",
    "Cross-functional Team Leadership",
    "Salesforce Education Cloud",
    "Agentic AI Development",
]


def iter_paragraphs(parent):
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    elif isinstance(parent, Table):
        parent_elm = parent._tbl
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            seen: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    tc_id = id(cell._tc)
                    if tc_id in seen:
                        continue
                    seen.add(tc_id)
                    yield from iter_paragraphs(cell)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def clone_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def clone_paragraph_after_from(after: Paragraph, source: Paragraph) -> Paragraph:
    new_p = deepcopy(source._p)
    after._p.addnext(new_p)
    return Paragraph(new_p, after._parent)


def set_run_bold(paragraph: Paragraph, bold: bool) -> None:
    for run in paragraph.runs:
        run.bold = bold


def paragraph_pstyle(paragraph: Paragraph) -> str | None:
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is None:
        return None
    style = p_pr.find(qn("w:pStyle"))
    if style is None:
        return None
    return style.get(qn("w:val"))


def contact_slot_token_inventory() -> set[str]:
    """Word placeholders for contact and links header."""
    return {
        "CONTACT_FIRST_NAME",
        "CONTACT_LAST_NAME",
        "CONTACT_EMAIL",
        "CONTACT_PHONE",
        "CONTACT_CITY",
        "CONTACT_STATE",
        "CONTACT_POSTAL_CODE",
        "CONTACT_LINKS_1_URL",
    }


def project_slot_token_inventory() -> set[str]:
    """Word placeholders for the single printed project slot."""
    return {"PROJECTS_1_TITLE", "PROJECTS_1_URL", "PROJECTS_1_SUMMARY"}


def right_column_section_markers(doc) -> list[str]:
    """Education -> divider -> Projects -> project-1 tokens -> divider -> Skills."""
    markers: list[str] = []
    started = False
    for para in iter_paragraphs(doc):
        text = para.text.strip()
        if text == "Education":
            started = True
            markers.append("Education")
            continue
        if not started:
            continue
        if paragraph_pstyle(para) == RIGHT_COLUMN_TOPBORDER_STYLE:
            markers.append("TOPBORDER")
        elif text == "Projects":
            markers.append("Projects")
        elif text == "{{PROJECTS_1_TITLE}}":
            markers.append("PROJECTS_1_TITLE")
        elif text == "{{PROJECTS_1_URL}}":
            markers.append("PROJECTS_1_URL")
        elif text == "{{PROJECTS_1_SUMMARY}}" or text == "{{PROJECTS_1_SUMMERY}}":
            markers.append("PROJECTS_1_SUMMARY")
        elif text == "Skills":
            markers.append("Skills")
            break
    return markers


def insert_projects_section(doc) -> None:
    """Insert Projects heading + project-1 tokens immediately before Skills.

    The sample already has a right-column top-border paragraph before Skills.
    Clone that divider in front of Projects so Education/Projects matches
    Projects/Skills (horizontal line plus the same top/bottom spacing).
    """
    education_heading: Paragraph | None = None
    last_edu_token: Paragraph | None = None
    skills_heading: Paragraph | None = None
    skills_topborder: Paragraph | None = None
    for para in iter_paragraphs(doc):
        text = para.text.strip()
        if text == "Education":
            education_heading = para
        elif text.startswith("{{EDUCATION_"):
            last_edu_token = para
        elif paragraph_pstyle(para) == RIGHT_COLUMN_TOPBORDER_STYLE:
            skills_topborder = para
        elif text == "Skills":
            skills_heading = para
            break
    if education_heading is None or last_edu_token is None or skills_heading is None:
        raise SystemExit(
            "cannot insert Projects: Education heading, education tokens, "
            "or Skills heading missing"
        )
    if skills_topborder is None:
        raise SystemExit(
            "cannot insert Projects: missing right-column top-border before Skills"
        )

    projects_heading = clone_paragraph_after_from(last_edu_token, education_heading)
    set_paragraph_text(projects_heading, "Projects")
    title_para = clone_paragraph_after_from(projects_heading, last_edu_token)
    set_paragraph_text(title_para, "{{PROJECTS_1_TITLE}}")
    set_run_bold(title_para, True)
    url_para = clone_paragraph_after_from(title_para, last_edu_token)
    set_paragraph_text(url_para, "{{PROJECTS_1_URL}}")
    set_run_bold(url_para, False)
    summary_para = clone_paragraph_after_from(url_para, last_edu_token)
    set_paragraph_text(summary_para, "{{PROJECTS_1_SUMMARY}}")
    set_run_bold(summary_para, False)
    clone_paragraph_after_from(last_edu_token, skills_topborder)


def job_slot_token_inventory(max_n: int = MAX_PRINTED_ROLES) -> set[str]:
    """Complete Word placeholder inventory for consecutive slots 1..max_n."""
    tokens: set[str] = set()
    for index in range(1, max_n + 1):
        tokens.update(
            {
                f"JOB_HISTORY_{index}_COMPANY",
                f"JOB_HISTORY_{index}_TITLE",
                f"JOB_HISTORY_{index}_LOCATION",
                f"JOB_HISTORY_{index}_START_DATE",
                f"JOB_HISTORY_{index}_END_DATE",
            }
        )
        for bullet in range(1, 6):
            tokens.add(f"JOB_HISTORY_{index}_B{bullet}")
    return tokens


def _collect_role_blocks(doc) -> dict[int, list[Paragraph]]:
    blocks: dict[int, list[Paragraph]] = {}
    for para in iter_paragraphs(doc):
        match = JOB_SLOT_TOKEN_RE.search(para.text)
        if match is None:
            continue
        blocks.setdefault(int(match.group(1)), []).append(para)
    return blocks


def remap_job_history_index(paragraph: Paragraph, src_n: int, dest_n: int) -> None:
    needle = f"JOB_HISTORY_{src_n}_"
    replacement = f"JOB_HISTORY_{dest_n}_"
    nodes = list(paragraph._p.iter(qn("w:t")))
    full = "".join(node.text or "" for node in nodes)
    if needle not in full:
        return
    updated = full.replace(needle, replacement)
    if not nodes:
        set_paragraph_text(paragraph, updated)
        return
    nodes[0].text = updated
    for node in nodes[1:]:
        node.text = ""


def extend_job_slots_to(doc, target: int = MAX_PRINTED_ROLES) -> None:
    """Clone the last complete styled role block (header + location + B1-B5)."""
    blocks = _collect_role_blocks(doc)
    if not blocks:
        raise SystemExit("no JOB_HISTORY slots found to extend")
    have = max(blocks)
    if have >= target:
        return
    source = blocks[have][:7]
    if len(source) < 7:
        raise SystemExit(
            f"role {have} is incomplete; need header, location, and B1-B5"
        )
    cursor = source[-1]
    for dest_n in range(have + 1, target + 1):
        for src_para in source:
            new_p = deepcopy(src_para._p)
            cursor._p.addnext(new_p)
            new_para = Paragraph(new_p, src_para._parent)
            remap_job_history_index(new_para, have, dest_n)
            cursor = new_para


def bind_role_headers_in_xml(docx_path: Path) -> None:
    """Replace job-header paragraphs that python-docx may not surface as a single run."""
    ET = __import__("xml.etree.ElementTree", fromlist=["ET"])
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    buffer = BytesIO(docx_path.read_bytes())
    with zipfile.ZipFile(buffer, "r") as zin:
        payload = {name: zin.read(name) for name in zin.namelist()}
    root = ET.fromstring(payload["word/document.xml"])
    for p_el in root.iter(f"{w}p"):
        nodes = list(p_el.iter(f"{w}t"))
        text = "".join(node.text or "" for node in nodes)
        for index, marker in enumerate(ROLE_HEADER_MARKERS, start=1):
            if marker in text and "{{JOB_HISTORY_" not in text:
                replacement = (
                    f"{{{{JOB_HISTORY_{index}_COMPANY}}}} - {{{{JOB_HISTORY_{index}_TITLE}}}}  "
                    f"{{{{JOB_HISTORY_{index}_START_DATE}}}} - {{{{JOB_HISTORY_{index}_END_DATE}}}}"
                )
                if nodes:
                    nodes[0].text = replacement
                    for node in nodes[1:]:
                        node.text = ""
                break
    payload["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in payload.items():
            zout.writestr(name, data)


def strip_resume_now_metadata(docx_path: Path) -> None:
    buffer = BytesIO(docx_path.read_bytes())
    with zipfile.ZipFile(buffer, "r") as zin:
        names = zin.namelist()
        payload = {name: zin.read(name) for name in names}

    payload.pop("docProps/custom.xml", None)

    rels_path = "_rels/.rels"
    if rels_path in payload:
        root = __import__("xml.etree.ElementTree", fromlist=["ET"]).fromstring(
            payload[rels_path]
        )
        ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        changed = False
        for rel in list(root.findall("r:Relationship", ns)):
            target = rel.get("Target", "")
            rel_type = rel.get("Type", "")
            if "custom.xml" in target or "custom-properties" in rel_type:
                root.remove(rel)
                changed = True
        if changed:
            payload[rels_path] = __import__(
                "xml.etree.ElementTree", fromlist=["ET"]
            ).tostring(root, encoding="utf-8", xml_declaration=True)

    ctypes_path = "[Content_Types].xml"
    if ctypes_path in payload:
        ET = __import__("xml.etree.ElementTree", fromlist=["ET"])
        root = ET.fromstring(payload[ctypes_path])
        ns = {"t": "http://schemas.openxmlformats.org/package/2006/content-types"}
        changed = False
        for node in list(root.findall("t:Override", ns)):
            if node.get("PartName") == "/docProps/custom.xml":
                root.remove(node)
                changed = True
        if changed:
            payload[ctypes_path] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True
            )

    if "docProps/core.xml" in payload:
        ET = __import__("xml.etree.ElementTree", fromlist=["ET"])
        root = ET.fromstring(payload["docProps/core.xml"])
        nsmap = {
            "dc": "http://purl.org/dc/elements/1.1/",
            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        }
        for tag in ("dc:title", "dc:creator", "cp:lastModifiedBy"):
            prefix, local = tag.split(":")
            for node in list(root.findall(f"{prefix}:{local}", nsmap)):
                node.text = ""
        payload["docProps/core.xml"] = ET.tostring(
            root, encoding="utf-8", xml_declaration=True
        )

    with zipfile.ZipFile(docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in payload.items():
            zout.writestr(name, data)


def unwrap_divider_tables(doc) -> None:
    """Replace nested 1-cell divider tables with direct divider paragraphs.

    In the sample, the dividers before Experience and Skills are wrapped in narrow
    nested tables (tblW=600 and tblW=660), causing short horizontal lines.
    Unwrapping them into direct paragraphs makes them span the full column width.
    """
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tc_elm = cell._tc
                for child in list(tc_elm):
                    if child.tag == qn("w:tbl"):
                        nested_tbl = Table(child, cell)
                        if len(nested_tbl.rows) == 1 and len(nested_tbl.rows[0].cells) == 1:
                            paras = list(nested_tbl.rows[0].cells[0].paragraphs)
                            if all(
                                paragraph_pstyle(p)
                                in {"topborder", RIGHT_COLUMN_TOPBORDER_STYLE}
                                for p in paras
                            ):
                                for p in paras:
                                    child.addprevious(deepcopy(p._p))
                                tc_elm.remove(child)


def derive() -> None:
    if not SAMPLE.is_file():
        raise SystemExit(f"missing layout sample: {SAMPLE}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SAMPLE, OUTPUT)

    doc = Document(str(OUTPUT))
    unwrap_divider_tables(doc)
    replacements: dict[str, str] = {
        "John Doe": "{{CONTACT_FIRST_NAME}} {{CONTACT_LAST_NAME}}",
        "john.doe@example.com / 555-555-5555 / New York, NY 10001": (
            "{{CONTACT_EMAIL}} | {{CONTACT_PHONE}} | {{CONTACT_CITY}}, "
            "{{CONTACT_STATE}} {{CONTACT_POSTAL_CODE}} | {{CONTACT_LINKS_1_URL}}"
        ),
        "Technical Architect and strategic leader with over 13 years of cross-industry expertise driving large-scale digital transformations and market expansions across the retail and education sectors. Skilled in applying structured frameworks (TOGAF) to develop reference architectures, design connected enterprise application ecosystems, and align technical execution with corporate business strategy.": "{{PROFESSIONAL_SUMMARY_SUMMARY}}",
        "05/2011": "{{EDUCATION_END_DATE}}",
        "University of Technology": "{{EDUCATION_INSTITUTION}}",
        "New York, NY": "{{EDUCATION_LOCATION}}",
        "B.S.: Information Technology Systems": "{{EDUCATION_DEGREE}}",
    }
    for index, bullets in enumerate(ROLE_BULLETS, start=1):
        for bullet_index, bullet in enumerate(bullets, start=1):
            replacements[bullet] = f"{{{{JOB_HISTORY_{index}_B{bullet_index}}}}}"
    for index, skill in enumerate(SAMPLE_SKILLS, start=1):
        replacements[skill] = f"{{{{SKILL_{index:02d}}}}}"

    location_hits = 0
    last_role_bullet: dict[int, Paragraph] = {}
    last_skill: Paragraph | None = None

    for para in iter_paragraphs(doc):
        text = para.text.strip()
        if text in ROLE_LOCATIONS:
            location_hits += 1
            if location_hits <= 5:
                set_paragraph_text(para, f"{{{{JOB_HISTORY_{location_hits}_LOCATION}}}}")
                continue
        header_bound = False
        for index, marker in enumerate(ROLE_HEADER_MARKERS, start=1):
            if marker in text and "{{" not in text:
                set_paragraph_text(
                    para,
                    f"{{{{JOB_HISTORY_{index}_COMPANY}}}} - {{{{JOB_HISTORY_{index}_TITLE}}}}  "
                    f"{{{{JOB_HISTORY_{index}_START_DATE}}}} - {{{{JOB_HISTORY_{index}_END_DATE}}}}",
                )
                header_bound = True
                break
        if header_bound:
            continue
        if text in replacements:
            set_paragraph_text(para, replacements[text])
            new_text = replacements[text]
            if new_text.startswith("{{JOB_HISTORY_") and "_B" in new_text:
                role_n = int(new_text.split("JOB_HISTORY_")[1].split("_")[0])
                last_role_bullet[role_n] = para
            if new_text.startswith("{{SKILL_"):
                last_skill = para

    for role_n, bullet_para in last_role_bullet.items():
        cursor = bullet_para
        for extra in (4, 5):
            cursor = clone_paragraph_after(cursor)
            set_paragraph_text(cursor, f"{{{{JOB_HISTORY_{role_n}_B{extra}}}}}")

    if last_skill is not None:
        cursor = last_skill
        for extra in range(31, 41):
            cursor = clone_paragraph_after(cursor)
            set_paragraph_text(cursor, f"{{{{SKILL_{extra:02d}}}}}")

    insert_projects_section(doc)
    extend_job_slots_to(doc, MAX_PRINTED_ROLES)
    doc.save(str(OUTPUT))
    bind_role_headers_in_xml(OUTPUT)
    strip_resume_now_metadata(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    derive()
