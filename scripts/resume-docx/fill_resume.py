#!/usr/bin/env python3
"""Fill locked-resume-template.docx from join-contract locked objects (ADR-015, ADR-021).

Discovers {{TOKENS}} in the template and resolves them by naming convention.
Does not read locked-job-search-criteria.json. Never writes .ai/guardrails/.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from lxml import etree

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_MODEL = REPO_ROOT / ".ai/history/architectural-decisions/resume-data-model.json"
CONTACT_PATH = REPO_ROOT / ".ai/guardrails/locked-contact.json"
EDUCATION_PATH = REPO_ROOT / ".ai/guardrails/locked-education.json"
SKILLS_PATH = REPO_ROOT / ".ai/guardrails/locked-skills.json"
JOB_HISTORY_PATH = REPO_ROOT / ".ai/guardrails/locked-job-history.json"
SUMMARY_PATH = REPO_ROOT / ".ai/guardrails/locked-professional-summary.json"
PROJECTS_PATH = REPO_ROOT / ".ai/guardrails/locked-projects.json"
TEMPLATE_PATH = REPO_ROOT / ".ai/guardrails/locked-resume-template.docx"
MASTER_RESUME = REPO_ROOT / ".ai/guardrails/locked-resume-master-template.md"
MIN_PRINTED_ROLES = 5
MAX_PRINTED_ROLES = 10
MIN_BULLETS_PER_ROLE = 1
MAX_BULLETS_PER_ROLE = 5
EXCLUDED_JOB_KEYS = ("11", "12")
JOB_BULLET_PLACEHOLDER_RE = re.compile(r"\{\{JOB_HISTORY_\d+_B\d+\}\}")
PROJECTS_TITLE_PLACEHOLDER_RE = re.compile(r"\{\{PROJECTS_\d+_TITLE\}\}")
PROJECTS_TOKEN_RE = re.compile(r"^PROJECTS_(\d+)_([A-Za-z0-9_]+)$", re.IGNORECASE)
RIGHT_COLUMN_TOPBORDER_STYLE = "divdocumentdivright-boxdisplaytabledisplaycelltopborder"

SECTION_FILES = {
    "contact": CONTACT_PATH,
    "education": EDUCATION_PATH,
    "skills": SKILLS_PATH,
    "job_history": JOB_HISTORY_PATH,
    "professional_summary": SUMMARY_PATH,
    "projects": PROJECTS_PATH,
}


def section_files() -> dict[str, Path]:
    return {
        "contact": CONTACT_PATH,
        "education": EDUCATION_PATH,
        "skills": SKILLS_PATH,
        "job_history": JOB_HISTORY_PATH,
        "professional_summary": SUMMARY_PATH,
        "projects": PROJECTS_PATH,
    }

TOKEN_RE = re.compile(r"\{\{([^{}]+)\}\}")
SKILL_TOKEN_RE = re.compile(r"^SKILL_(\d+)$", re.IGNORECASE)
JOB_BULLET_TOKEN_RE = re.compile(r"^JOB_HISTORY_(\d+)_B(\d+)$", re.IGNORECASE)
JOB_HEADER_RE = re.compile(
    r"^\{\{JOB_HISTORY_(\d+)_COMPANY\}\} - \{\{JOB_HISTORY_(\d+)_TITLE\}\}\s+"
    r"\{\{JOB_HISTORY_(\d+)_START_DATE\}\} - \{\{JOB_HISTORY_(\d+)_END_DATE\}\}\s*$"
)
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
# Palatino Linotype 10pt advance widths in twips (Windows GDI @ 96 DPI).
# Job headers use Strong1 (bold). Unknown glyphs use bold "n".
PALATINO_10PT_BOLD_TWIPS = {
    " ": 45.0, "a": 105.0, "b": 120.0, "c": 90.0, "d": 120.0, "e": 105.0,
    "f": 75.0, "g": 120.0, "h": 120.0, "i": 60.0, "j": 60.0, "k": 120.0,
    "l": 60.0, "m": 180.0, "n": 120.0, "o": 120.0, "p": 120.0, "q": 120.0,
    "r": 75.0, "s": 90.0, "t": 60.0, "u": 120.0, "v": 105.0, "w": 165.0,
    "x": 90.0, "y": 105.0, "z": 105.0, "A": 150.0, "B": 135.0, "C": 135.0,
    "D": 165.0, "E": 120.0, "F": 120.0, "G": 165.0, "H": 165.0, "I": 75.0,
    "J": 75.0, "K": 150.0, "L": 120.0, "M": 195.0, "N": 165.0, "O": 165.0,
    "P": 120.0, "Q": 165.0, "R": 135.0, "S": 120.0, "T": 135.0, "U": 150.0,
    "V": 150.0, "W": 195.0, "X": 135.0, "Y": 135.0, "Z": 135.0, "0": 105.0,
    "1": 105.0, "2": 105.0, "3": 105.0, "4": 105.0, "5": 105.0, "6": 105.0,
    "7": 105.0, "8": 105.0, "9": 105.0, "/": 60.0, "-": 60.0, ".": 60.0,
    ",": 60.0, "'": 45.0, "&": 165.0, "(": 60.0, ")": 60.0, "[": 60.0,
    "]": 60.0,
}
DATE_COL_PAD_TWIPS = 90.0
MIN_DATE_COL_TWIPS = 1200.0


class FillError(Exception):
    """Human-readable fill abort."""


def load_json(path: Path) -> dict:
    if not path.is_file():
        raise FillError(f"missing required file: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise FillError(f"{path}: invalid JSON — {exc}") from exc


def require_fields(obj: dict, fields: list[str], label: str) -> None:
    missing = [field for field in fields if field not in obj or obj[field] in (None, "")]
    if missing:
        raise FillError(f"{label} missing required field(s): {', '.join(missing)}")


def is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def iter_projects_entries(projects: dict) -> list[tuple[str, dict]]:
    entries: list[tuple[str, dict]] = []
    for key, value in projects.items():
        if key == "id":
            continue
        if not str(key).isdigit():
            raise FillError(f"projects has non-digit key {key!r}")
        if not isinstance(value, dict):
            raise FillError(f"projects[{key}] is not an object")
        entries.append((str(key), value))
    return entries


def validate_projects_entries(projects: dict) -> None:
    require_fields(projects, ["id"], "projects")
    for key, entry in iter_projects_entries(projects):
        missing = [field for field in ("title", "summary") if field not in entry]
        if missing:
            raise FillError(f"projects[{key}] missing required field(s): {', '.join(missing)}")
        title_blank = is_blank(entry.get("title"))
        summary_blank = is_blank(entry.get("summary"))
        if title_blank != summary_blank:
            raise FillError(
                f"projects[{key}] must populate both title and summary or neither"
            )


def project_1_populated(projects: dict) -> bool:
    entry = projects.get("1")
    if not isinstance(entry, dict):
        return False
    return not is_blank(entry.get("title")) and not is_blank(entry.get("summary"))


def parse_tailored_resume(path: Path) -> tuple[list[dict], list[str]]:
    if not path.is_file():
        raise FillError(f"missing required file: {path}")
    roles: list[dict] = []
    skills: list[str] = []
    current: dict | None = None
    in_skills = False
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("## "):
            header = stripped[3:].strip()
            if header.lower() == "skills":
                if current is not None:
                    roles.append(current)
                    current = None
                in_skills = True
                continue
            if current is not None:
                roles.append(current)
            current = {"title": header, "bullets": []}
            in_skills = False
            continue
        if stripped.startswith("- "):
            item = stripped[2:].strip()
            item = re.sub(r"^\*\*(.+?)\*\*\s*", r"\1 ", item).strip()
            if in_skills:
                skills.append(item)
            elif current is not None:
                current["bullets"].append(item)
    if current is not None:
        roles.append(current)
    if not roles:
        raise FillError(f"{path}: no role sections found")
    for role in roles:
        count = len(role["bullets"])
        if count < MIN_BULLETS_PER_ROLE or count > MAX_BULLETS_PER_ROLE:
            raise FillError(
                f"{path}: role '{role['title']}' has {count} bullets; "
                f"required {MIN_BULLETS_PER_ROLE}-{MAX_BULLETS_PER_ROLE}"
            )
        for index, bullet in enumerate(role["bullets"], start=1):
            colon = bullet.find(":")
            theme = bullet[:colon].strip() if colon >= 0 else ""
            if colon < 0 or not theme:
                raise FillError(
                    f"{path}: role '{role['title']}' bullet {index} is missing "
                    "a theme before the first colon"
                )
    return roles, skills


def parse_role_headers(path: Path) -> list[str]:
    headers: list[str] = []
    if not path.is_file():
        raise FillError(f"missing required file: {path}")
    for raw in path.read_text(encoding="utf-8").splitlines():
        stripped = raw.strip()
        if not stripped.startswith("## "):
            continue
        header = stripped[3:].strip()
        if header.lower() == "skills":
            break
        headers.append(header)
    return headers


def excluded_resume_titles(job_history: dict) -> set[str]:
    titles: set[str] = set()
    for key in EXCLUDED_JOB_KEYS:
        role = job_history.get(key)
        if isinstance(role, dict) and role.get("title"):
            titles.add(str(role["title"]).casefold())
    return titles


def validate_join(model: dict, bundles: dict[str, dict]) -> None:
    contact = bundles["contact"]
    require_fields(contact, model["contact"]["required_fields"], "contact")
    require_fields(bundles["education"], model["education"]["required_fields"], "education")
    require_fields(bundles["skills"], model["skills"]["required_fields"], "skills")
    require_fields(
        bundles["professional_summary"],
        model["professional_summary"]["required_fields"],
        "professional_summary",
    )
    require_fields(
        bundles["job_history"], model["job_history"]["required_fields"], "job_history"
    )
    require_fields(bundles["projects"], model["projects"]["required_fields"], "projects")
    validate_projects_entries(bundles["projects"])
    expected = {
        "education_id": bundles["education"]["id"],
        "skills_id": bundles["skills"]["id"],
        "job_history_id": bundles["job_history"]["id"],
        "professional_summary_id": bundles["professional_summary"]["id"],
        "projects_id": bundles["projects"]["id"],
    }
    mismatches = [
        f"contact.{fk}={contact[fk]!r} != {child_id!r}"
        for fk, child_id in expected.items()
        if contact.get(fk) != child_id
    ]
    if mismatches:
        raise FillError("foreign key mismatch: " + "; ".join(mismatches))


def tailored_roles_by_prefix(
    tailored_roles: list[dict],
    job_history: dict,
    master_headers: list[str],
) -> dict[str, dict]:
    count = len(tailored_roles)
    if count < MIN_PRINTED_ROLES or count > MAX_PRINTED_ROLES:
        raise FillError(
            f"tailored-resume.md has {count} role sections; "
            f"required consecutive prefix {MIN_PRINTED_ROLES}-{MAX_PRINTED_ROLES}"
        )
    if len(master_headers) < count:
        raise FillError(
            "locked-resume-master-template.md has "
            f"{len(master_headers)} roles; need at least {count} for this prefix"
        )
    expected = master_headers[:count]
    actual = [role["title"] for role in tailored_roles]
    if actual != expected:
        raise FillError(
            "role sections must match the consecutive master-template prefix "
            f"1..{count} with no skipping or reordering.\n"
            f"  expected: {expected}\n"
            f"  actual:   {actual}"
        )
    excluded = excluded_resume_titles(job_history)
    blocked = [title for title in actual if title.casefold() in excluded]
    if blocked:
        raise FillError(
            "locked keys 11-12 must not appear on the resume: "
            + ", ".join(repr(title) for title in blocked)
        )
    matched: dict[str, dict] = {}
    for index, tailored in enumerate(tailored_roles, start=1):
        key = str(index)
        payload = job_history.get(key)
        if not isinstance(payload, dict):
            raise FillError(f"job_history[{key}] is missing")
        missing = [
            field
            for field in ("title", "company", "location", "start_date", "end_date")
            if not payload.get(field)
        ]
        if missing:
            raise FillError(f"job_history[{key}] missing {', '.join(missing)}")
        matched[key] = {
            **payload,
            "title": tailored["title"],
            "bullets": tailored["bullets"],
        }
    return matched


def iter_paragraphs(parent):
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    elif isinstance(parent, Table):
        parent_elm = parent._tbl
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            seen: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    tc_id = id(cell._tc)
                    if tc_id in seen:
                        continue
                    seen.add(tc_id)
                    yield from iter_paragraphs(cell)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        _preserve_w_t_spaces(paragraph.runs[0])
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def _preserve_w_t_spaces(run) -> None:
    """Keep consecutive spaces inside <w:t>; python-docx only flags leading/trailing."""
    for node in run._r.iter(qn("w:t")):
        if node.text and "  " in node.text:
            node.set(XML_SPACE, "preserve")


def _copy_rpr(r_pr):
    return deepcopy(r_pr) if r_pr is not None else None


def _run_r_pr(run):
    r_pr = run._r.find(qn("w:rPr"))
    return _copy_rpr(r_pr)


def _rpr_key(r_pr) -> str:
    if r_pr is None:
        return ""
    xml = getattr(r_pr, "xml", None)
    if xml is not None:
        return str(xml)
    return etree.tostring(r_pr, encoding="unicode")


def _text_runs(paragraph: Paragraph) -> list[Run]:
    """Runs whose text belongs to this paragraph, including hyperlink children."""
    p_el = paragraph._p
    runs: list[Run] = []
    for r_el in p_el.iter(qn("w:r")):
        parent_p = r_el
        while parent_p is not None and parent_p.tag != qn("w:p"):
            parent_p = parent_p.getparent()
        if parent_p is p_el:
            runs.append(Run(r_el, paragraph))
    return runs


def paragraph_run_parts(paragraph: Paragraph) -> list[tuple[str, object]]:
    return [(run.text or "", _run_r_pr(run)) for run in _text_runs(paragraph)]


def _majority_part_index(owners: list[int], start: int, end: int, parts: list[tuple[str, object]]) -> int:
    """Pick the run whose rPr covers the most characters of a split token."""
    rpr_count: dict[str, int] = {}
    rpr_first_idx: dict[str, int] = {}
    for idx in owners[start:end]:
        key = _rpr_key(parts[idx][1])
        rpr_count[key] = rpr_count.get(key, 0) + 1
        rpr_first_idx.setdefault(key, idx)
    best_key = max(rpr_count, key=lambda key: (rpr_count[key], -rpr_first_idx[key]))
    return rpr_first_idx[best_key]


def replace_tokens_in_parts(
    parts: list[tuple[str, object]], mapping: dict[str, str]
) -> list[tuple[str, object]]:
    """Replace {{TOKENS}} without collapsing mixed run formatting on the same line.

    Word often splits one placeholder across runs. Opening braces may keep the
    line's starting style while the token body has italic/unbold. The filled
    value uses the rPr that covers the most characters of that token.
    """
    if not parts:
        return parts
    full = "".join(text for text, _ in parts)
    if "{{" not in full:
        return parts
    owners: list[int] = []
    for index, (text, _) in enumerate(parts):
        owners.extend([index] * len(text))
    spans: list[tuple[int, int, str]] = []
    for match in TOKEN_RE.finditer(full):
        key = "{{" + match.group(1) + "}}"
        if key in mapping:
            spans.append((match.start(), match.end(), mapping[key]))
    if not spans:
        return parts
    span_at = {start: (end, value) for start, end, value in spans}
    buckets: list[list[str]] = [[] for _ in parts]
    index = 0
    length = len(full)
    while index < length:
        if index in span_at:
            end, value = span_at[index]
            winner = _majority_part_index(owners, index, end, parts)
            buckets[winner].append(value)
            index = end
            continue
        buckets[owners[index]].append(full[index])
        index += 1
    return [("".join(buckets[i]), parts[i][1]) for i in range(len(parts))]


def split_parts_at(
    parts: list[tuple[str, object]], char_index: int
) -> tuple[list[tuple[str, object]], list[tuple[str, object]]]:
    left: list[tuple[str, object]] = []
    right: list[tuple[str, object]] = []
    pos = 0
    for text, r_pr in parts:
        end = pos + len(text)
        if end <= char_index:
            left.append((text, r_pr))
        elif pos >= char_index:
            right.append((text, r_pr))
        else:
            cut = char_index - pos
            if cut > 0:
                left.append((text[:cut], _copy_rpr(r_pr)))
            if cut < len(text):
                right.append((text[cut:], _copy_rpr(r_pr)))
        pos = end
    return left, right


def apply_parts_to_paragraph(paragraph: Paragraph, parts: list[tuple[str, object]]) -> None:
    runs = _text_runs(paragraph)
    for run, (text, _) in zip(runs, parts):
        run.text = text
        _preserve_w_t_spaces(run)
    if len(runs) < len(parts):
        raise FillError("placeholder run rewrite lost template runs")


def _rpr_sz_pt(r_element) -> float | None:
    r_pr = r_element.find(qn("w:rPr"))
    if r_pr is None:
        return None
    sz = r_pr.find(qn("w:sz"))
    if sz is None:
        return None
    raw = sz.get(qn("w:val"))
    if not raw:
        return None
    try:
        return int(raw) / 2.0
    except ValueError:
        return None


def paragraph_font_pt(paragraph: Paragraph) -> float:
    for run in paragraph.runs:
        size = run.font.size
        if size is not None:
            return float(size.pt)
        from_rpr = _rpr_sz_pt(run._r)
        if from_rpr is not None:
            return from_rpr
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        r_pr = p_pr.find(qn("w:rPr"))
        if r_pr is not None:
            sz = r_pr.find(qn("w:sz"))
            if sz is not None and sz.get(qn("w:val")):
                try:
                    return int(sz.get(qn("w:val"))) / 2.0
                except ValueError:
                    pass
    return 10.0


def text_width_twips(text: str, size_pt: float) -> float:
    scale = size_pt / 10.0
    fallback = PALATINO_10PT_BOLD_TWIPS["n"]
    return sum(PALATINO_10PT_BOLD_TWIPS.get(ch, fallback) for ch in text) * scale


def containing_cell_width_twips(paragraph: Paragraph) -> float | None:
    node = paragraph._p.getparent()
    while node is not None:
        if node.tag == qn("w:tc"):
            tc_pr = node.find(qn("w:tcPr"))
            if tc_pr is None:
                return None
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                return None
            width_type = tc_w.get(qn("w:type"))
            raw = tc_w.get(qn("w:w"))
            if not raw or width_type not in (None, "dxa"):
                return None
            try:
                return float(raw)
            except ValueError:
                return None
        node = node.getparent()
    return None


def _oxml(tag: str):
    return OxmlElement(tag)


def _set_dxa(element, width_twips: float) -> None:
    element.set(qn("w:w"), str(max(1, int(round(width_twips)))))
    element.set(qn("w:type"), "dxa")


def _nil_borders() -> object:
    borders = _oxml("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = _oxml(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    return borders


def _zero_cell_margins() -> object:
    mar = _oxml("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = _oxml(f"w:{edge}")
        _set_dxa(el, 0)
        mar.append(el)
    return mar


def _header_p_pr(paragraph: Paragraph, align_right: bool):
    p_pr = paragraph._p.find(qn("w:pPr"))
    cloned = deepcopy(p_pr) if p_pr is not None else _oxml("w:pPr")
    for child in list(cloned):
        if child.tag in {qn("w:tabs"), qn("w:jc"), qn("w:rPr")}:
            cloned.remove(child)
    if align_right:
        jc = _oxml("w:jc")
        jc.set(qn("w:val"), "right")
        cloned.append(jc)
    return cloned


def _header_paragraph_element(parts: list[tuple[str, object]], p_pr):
    p_el = _oxml("w:p")
    p_el.append(deepcopy(p_pr))
    wrote = False
    for text, r_pr in parts:
        if text == "":
            continue
        run = _oxml("w:r")
        if r_pr is not None:
            run.append(deepcopy(r_pr))
        t_el = _oxml("w:t")
        t_el.text = text
        if text[:1] in " \t" or text[-1:] in " \t" or "  " in text:
            t_el.set(XML_SPACE, "preserve")
        run.append(t_el)
        p_el.append(run)
        wrote = True
    if not wrote:
        p_el.append(_oxml("w:r"))
    return p_el


def _header_cell(
    width_twips: float, parts: list[tuple[str, object]], p_pr, *, nowrap: bool
) -> object:
    cell = _oxml("w:tc")
    tc_pr = _oxml("w:tcPr")
    tc_w = _oxml("w:tcW")
    _set_dxa(tc_w, width_twips)
    tc_pr.append(tc_w)
    if nowrap:
        tc_pr.append(_oxml("w:noWrap"))
        v_align = _oxml("w:vAlign")
        v_align.set(qn("w:val"), "top")
        tc_pr.append(v_align)
    cell.append(tc_pr)
    cell.append(_header_paragraph_element(parts, p_pr))
    return cell


def replace_job_header_with_split_row(
    paragraph: Paragraph,
    left_parts: list[tuple[str, object]],
    right_parts: list[tuple[str, object]],
) -> None:
    """Put dates in a right cell so they stay on the header row as title length changes.

    Space-padding a single paragraph wraps in this template: Palatino 10pt
    company+title+dates is wider than the 6720-twip experience column.
    Each side keeps the template run formatting for its placeholders.
    """
    right = "".join(text for text, _ in right_parts)
    avail = containing_cell_width_twips(paragraph) or 6720.0
    size_pt = paragraph_font_pt(paragraph)
    date_w = max(
        MIN_DATE_COL_TWIPS,
        text_width_twips(right, size_pt) + DATE_COL_PAD_TWIPS,
    )
    if date_w >= avail - 600:
        date_w = max(MIN_DATE_COL_TWIPS, avail * 0.28)
    left_w = max(600.0, avail - date_w)
    p_pr_left = _header_p_pr(paragraph, align_right=False)
    p_pr_right = _header_p_pr(paragraph, align_right=True)
    tbl = _oxml("w:tbl")
    tbl_pr = _oxml("w:tblPr")
    tbl_w = _oxml("w:tblW")
    _set_dxa(tbl_w, avail)
    tbl_pr.append(tbl_w)
    layout = _oxml("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_pr.append(_nil_borders())
    tbl_pr.append(_zero_cell_margins())
    tbl.append(tbl_pr)
    grid = _oxml("w:tblGrid")
    for width in (left_w, date_w):
        col = _oxml("w:gridCol")
        col.set(qn("w:w"), str(int(round(width))))
        grid.append(col)
    tbl.append(grid)
    row = _oxml("w:tr")
    row.append(_header_cell(left_w, left_parts, p_pr_left, nowrap=False))
    row.append(_header_cell(date_w, right_parts, p_pr_right, nowrap=True))
    tbl.append(row)
    paragraph._p.addnext(tbl)
    remove_paragraph(paragraph)


def apply_job_header(paragraph: Paragraph, mapping: dict[str, str], match: re.Match) -> bool:
    company_key, title_key, start_key, end_key = match.groups()
    if not (company_key == title_key == start_key == end_key):
        replace_placeholders(paragraph, mapping)
        return False
    company = mapping.get(f"{{{{JOB_HISTORY_{company_key}_COMPANY}}}}", "")
    title = mapping.get(f"{{{{JOB_HISTORY_{title_key}_TITLE}}}}", "")
    start = mapping.get(f"{{{{JOB_HISTORY_{start_key}_START_DATE}}}}", "")
    end = mapping.get(f"{{{{JOB_HISTORY_{end_key}_END_DATE}}}}", "")
    if not any((company, title, start, end)):
        set_paragraph_text(paragraph, "")
        return False
    parts = paragraph_run_parts(paragraph)
    full = "".join(text for text, _ in parts)
    split_at = full.find(f"{{{{JOB_HISTORY_{start_key}_START_DATE}}}}")
    if split_at < 0:
        replace_placeholders(paragraph, mapping)
        return False
    left_parts, right_parts = split_parts_at(parts, split_at)
    replace_job_header_with_split_row(
        paragraph,
        replace_tokens_in_parts(left_parts, mapping),
        replace_tokens_in_parts(right_parts, mapping),
    )
    return True


def replace_placeholders(paragraph: Paragraph, mapping: dict[str, str]) -> None:
    text = paragraph.text
    if "{{" not in text:
        return
    parts = replace_tokens_in_parts(paragraph_run_parts(paragraph), mapping)
    filled = "".join(chunk for chunk, _ in parts)
    if filled != text:
        apply_parts_to_paragraph(paragraph, parts)


def _rpr_with_bold(r_pr, bold: bool):
    cloned = deepcopy(r_pr) if r_pr is not None else _oxml("w:rPr")
    for child in list(cloned):
        if child.tag in {qn("w:b"), qn("w:bCs")}:
            cloned.remove(child)
    bold_el = _oxml("w:b")
    if not bold:
        bold_el.set(qn("w:val"), "0")
    cloned.append(bold_el)
    return cloned


def _replace_paragraph_runs(
    paragraph: Paragraph, parts: list[tuple[str, object]]
) -> None:
    p_el = paragraph._p
    for child in list(p_el):
        if child.tag in {qn("w:r"), qn("w:hyperlink")}:
            p_el.remove(child)
    for text, r_pr in parts:
        run = _oxml("w:r")
        if r_pr is not None:
            run.append(deepcopy(r_pr))
        t_el = _oxml("w:t")
        t_el.text = text
        if text[:1] in " \t" or text[-1:] in " \t" or "  " in text:
            t_el.set(XML_SPACE, "preserve")
        run.append(t_el)
        p_el.append(run)


def apply_theme_colon_bold(paragraph: Paragraph) -> None:
    """Bold from the start through the first colon (colon included)."""
    parts = paragraph_run_parts(paragraph)
    full = "".join(text for text, _ in parts)
    colon = full.find(":")
    if colon < 0:
        return
    owners: list[int] = []
    for index, (text, _) in enumerate(parts):
        owners.extend([index] * len(text))
    theme_rpr = parts[owners[0]][1] if owners else None
    body_index = owners[colon + 1] if colon + 1 < len(owners) else owners[0]
    body_rpr = parts[body_index][1] if parts else None
    theme_parts = [(full[: colon + 1], _rpr_with_bold(theme_rpr, True))]
    body = full[colon + 1 :]
    if body:
        theme_parts.append((body, _rpr_with_bold(body_rpr, False)))
    _replace_paragraph_runs(paragraph, theme_parts)


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._p
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_run_bold(paragraph: Paragraph, bold: bool) -> None:
    for run in paragraph.runs:
        run.bold = bold


def paragraph_pstyle(paragraph: Paragraph) -> str | None:
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is None:
        return None
    style = p_pr.find(qn("w:pStyle"))
    if style is None:
        return None
    return style.get(qn("w:val"))


def company_title_case(company_slug: str) -> str:
    return "_".join(part.capitalize() for part in company_slug.split("-"))


def output_name(sandbox: Path) -> str:
    company_slug = sandbox.parent.name
    return f"John_Doe_Resume_{company_title_case(company_slug)}.docx"


def collect_template_tokens(doc) -> set[str]:
    tokens: set[str] = set()
    for para in iter_paragraphs(doc):
        tokens.update(TOKEN_RE.findall(para.text))
    return tokens


def stringify_value(value, token: str) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float)) and not isinstance(value, bool):
        return str(value)
    raise FillError(f"unresolved placeholder {token}: non-scalar value")


def walk_value(current, parts: list[str], token: str):
    if not parts:
        return current
    if isinstance(current, list):
        if not parts[0].isdigit():
            raise FillError(f"unresolved placeholder {token}: expected array index")
        index = int(parts[0])
        if index < 1 or index > len(current):
            return ""
        return walk_value(current[index - 1], parts[1:], token)
    if isinstance(current, dict):
        if parts[0].isdigit():
            key = parts[0]
            if key not in current:
                return ""
            return walk_value(current[key], parts[1:], token)
        for length in range(len(parts), 0, -1):
            field = "_".join(parts[:length]).lower()
            if field in current:
                return walk_value(current[field], parts[length:], token)
        raise FillError(
            f"unknown field {'_'.join(parts).lower()} in placeholder {token}"
        )
    raise FillError(f"unresolved placeholder {token}: cannot walk into scalar")


def match_object_prefix(body: str, object_names: list[str]) -> tuple[str, str]:
    upper = body.upper()
    for name in sorted(object_names, key=len, reverse=True):
        prefix = name.upper()
        if upper == prefix:
            raise FillError(f"unresolved placeholder {{{{{body}}}}}: object with no field")
        if upper.startswith(prefix + "_"):
            return name, body[len(prefix) + 1 :]
    raise FillError(f"unknown placeholder: {{{{{body}}}}}")


def resolve_token(
    body: str,
    bundles: dict[str, dict],
    tailored_by_key: dict[str, dict],
    skills: list[str],
) -> str:
    token = "{{" + body + "}}"
    skill_match = SKILL_TOKEN_RE.fullmatch(body.strip())
    if skill_match:
        index = int(skill_match.group(1))
        if index < 1:
            raise FillError(f"unknown placeholder: {token}")
        if index > len(skills):
            return ""
        return skills[index - 1]

    bullet_match = JOB_BULLET_TOKEN_RE.fullmatch(body.strip())
    if bullet_match:
        key = bullet_match.group(1)
        bullet_n = int(bullet_match.group(2))
        if bullet_n < 1:
            raise FillError(f"unknown placeholder: {token}")
        role = tailored_by_key.get(key)
        if role is None:
            return ""
        bullets = role["bullets"]
        if bullet_n > len(bullets):
            return ""
        return bullets[bullet_n - 1]

    projects_match = PROJECTS_TOKEN_RE.fullmatch(body.strip())
    if projects_match:
        key = projects_match.group(1)
        field = projects_match.group(2).lower()
        if field == "summery":
            field = "summary"
        entry = bundles.get("projects", {}).get(key)
        if not isinstance(entry, dict):
            return ""
        return stringify_value(entry.get(field), token)

    object_name, remainder = match_object_prefix(body.strip(), list(bundles.keys()))
    parts = [part for part in remainder.split("_") if part]
    if not parts:
        raise FillError(f"unresolved placeholder {token}: object with no field")

    if object_name == "job_history":
        if not parts[0].isdigit():
            raise FillError(f"unknown field {remainder.lower()} in placeholder {token}")
        key = parts[0]
        locked_role = bundles["job_history"].get(key)
        if key not in tailored_by_key:
            if locked_role is None:
                return ""
            stringify_value(walk_value(locked_role, parts[1:], token), token)
            return ""
        if locked_role is None:
            return ""
        return stringify_value(walk_value(locked_role, parts[1:], token), token)

    return stringify_value(walk_value(bundles[object_name], parts, token), token)


def build_mapping(
    tokens: set[str],
    bundles: dict[str, dict],
    tailored_by_key: dict[str, dict],
    skills: list[str],
) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for body in sorted(tokens):
        mapping["{{" + body + "}}"] = resolve_token(
            body, bundles, tailored_by_key, skills
        )
    return mapping


def fill_document_clean(
    template: Path,
    mapping: dict[str, str],
    dest: Path,
    remove_projects_section: bool = False,
) -> None:
    shutil.copy2(template, dest)
    doc = Document(str(dest))
    to_remove: list[Paragraph] = []
    paragraphs = list(iter_paragraphs(doc))
    for index, para in enumerate(paragraphs):
        before = para.text
        if remove_projects_section and before.strip() == "Projects":
            prior = paragraphs[index - 1] if index else None
            if prior is not None and paragraph_pstyle(prior) == RIGHT_COLUMN_TOPBORDER_STYLE:
                to_remove.append(prior)
            to_remove.append(para)
            continue
        if "{{" not in before:
            continue
        is_job_bullet = bool(JOB_BULLET_PLACEHOLDER_RE.search(before))
        is_projects_title = bool(PROJECTS_TITLE_PLACEHOLDER_RE.search(before))
        header_match = JOB_HEADER_RE.match(before.strip())
        if header_match:
            if apply_job_header(para, mapping, header_match):
                continue
        else:
            replace_placeholders(para, mapping)
        after = para.text
        if "{{" in after and "}}" in after:
            raise FillError(f"unresolved placeholder remaining: {after}")
        if after.strip() == "":
            to_remove.append(para)
        elif is_job_bullet:
            apply_theme_colon_bold(para)
        elif is_projects_title:
            set_run_bold(para, True)
    for para in to_remove:
        remove_paragraph(para)
    doc.save(str(dest))


def gather_and_fill(sandbox: Path) -> Path:
    model = load_json(DATA_MODEL)
    bundles = {name: load_json(path) for name, path in section_files().items()}
    validate_join(model, bundles)
    if not TEMPLATE_PATH.is_file():
        raise FillError(f"missing required file: {TEMPLATE_PATH}")
    tailored_path = sandbox / "tailored-resume.md"
    tailored_roles, skills = parse_tailored_resume(tailored_path)
    master_headers = parse_role_headers(MASTER_RESUME)
    tailored_by_key = tailored_roles_by_prefix(
        tailored_roles, bundles["job_history"], master_headers
    )
    template_doc = Document(str(TEMPLATE_PATH))
    tokens = collect_template_tokens(template_doc)
    mapping = build_mapping(tokens, bundles, tailored_by_key, skills)
    dest = sandbox / output_name(sandbox)
    try:
        fill_document_clean(
            TEMPLATE_PATH,
            mapping,
            dest,
            remove_projects_section=not project_1_populated(bundles["projects"]),
        )
    except PermissionError as exc:
        raise FillError(
            f"cannot write {dest} - close the file if it is open in Word"
        ) from exc
    return dest


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Fill locked-resume-template.docx for a resume-tailor sandbox."
    )
    parser.add_argument(
        "--sandbox",
        required=True,
        help="Path to .ai/history/resume-tailor/{company}/{position}",
    )
    args = parser.parse_args()
    sandbox = Path(args.sandbox)
    if not sandbox.is_dir():
        print(f"ERROR: sandbox directory not found: {sandbox}", file=sys.stderr)
        return 1
    try:
        dest = gather_and_fill(sandbox)
    except FillError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"wrote {dest}")
    _print_fill_estimate(dest)
    return 0


def _print_fill_estimate(dest: Path) -> None:
    """Print the advisory last-page fill estimate (ADR-051).

    Advisory only: a failure here must never fail the fill. The accordion
    page-fill re-evaluation (resume-tailor SKILL.md) reads this line when N>5.
    """
    try:
        import page_fill

        print(page_fill.estimate_fill(dest).summary_line())
    except Exception as exc:  # noqa: BLE001 - advisory, never fatal
        print(f"note: last-page fill estimate unavailable ({exc})")


if __name__ == "__main__":
    sys.exit(main())
