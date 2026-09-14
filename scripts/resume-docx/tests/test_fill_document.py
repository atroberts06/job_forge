"""Word fill: theme bolding and unused slot/paragraph removal."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

_DOCX_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_DOCX_DIR))

import fill_resume as filler
from test_fill_validation import MASTER_TITLES, _resume, _themed


def _sandbox(tmp_path: Path, titles: list[str], bullets: list[list[str]]) -> Path:
    sandbox = tmp_path / "acme" / "architect"
    sandbox.mkdir(parents=True)
    _resume(sandbox / "tailored-resume.md", titles, bullets)
    return sandbox


def _job_paragraphs(docx_path: Path) -> list[str]:
    texts: list[str] = []
    for para in filler.iter_paragraphs(Document(str(docx_path))):
        text = para.text.strip()
        if text:
            texts.append(text)
    return texts


def _right_column_markers(docx_path: Path) -> list[str]:
    markers: list[str] = []
    started = False
    for para in filler.iter_paragraphs(Document(str(docx_path))):
        text = para.text.strip()
        if text == "Education":
            started = True
            markers.append("Education")
            continue
        if not started:
            continue
        if filler.paragraph_pstyle(para) == filler.RIGHT_COLUMN_TOPBORDER_STYLE:
            markers.append("TOPBORDER")
        elif text:
            markers.append(text)
        if text == "Skills":
            break
    return markers


def test_n5_success_removes_unused_slots_and_optional_bullets(tmp_path: Path) -> None:
    sandbox = _sandbox(
        tmp_path,
        MASTER_TITLES[:5],
        [_themed(3), _themed(3), _themed(3), _themed(3), _themed(1)],
    )
    dest = filler.gather_and_fill(sandbox)
    texts = _job_paragraphs(dest)
    joined = "\n".join(texts)
    assert "Northstar Retail" not in joined
    assert "Application Development Systems Analyst" not in joined
    assert "{{" not in joined
    assert any("Theme 1: Locked result 1." in text for text in texts)
    assert not any("Theme 4:" in text for text in texts)
    companies = [text for text in texts if "Enterprise Solutions Inc. - Team Lead" in text]
    assert len(companies) == 2


def test_empty_projects_removes_entire_section(tmp_path: Path, monkeypatch) -> None:
    proj_path = tmp_path / "locked-projects.json"
    proj_path.write_text(
        json.dumps(
            {
                "id": "prj_001",
                "1": {
                    "title": None,
                    "summary": None,
                },
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(filler, "PROJECTS_PATH", proj_path)
    sandbox = _sandbox(
        tmp_path,
        MASTER_TITLES[:5],
        [_themed(1) for _ in MASTER_TITLES[:5]],
    )
    dest = filler.gather_and_fill(sandbox)
    texts = _job_paragraphs(dest)
    joined = "\n".join(texts)
    assert "Projects" not in texts
    assert "{{PROJECTS_" not in joined
    edu = texts.index("Education")
    skills = texts.index("Skills")
    assert edu < skills
    assert "Projects" not in texts[edu:skills]
    markers = _right_column_markers(dest)
    assert markers.count("TOPBORDER") == 1
    assert markers.index("Education") < markers.index("TOPBORDER") < markers.index("Skills")


def test_populated_project_renders_title_summary_and_hides_slot_two(
    tmp_path: Path, monkeypatch
) -> None:
    proj_path = tmp_path / "locked-projects.json"
    proj_path.write_text(
        json.dumps(
            {
                "id": "prj_001",
                "1": {
                    "title": "Career Forge",
                    "summary": "Automated job-search pipeline.",
                },
                "2": {
                    "title": "Hidden Project",
                    "summary": "Must not print.",
                },
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(filler, "PROJECTS_PATH", proj_path)
    sandbox = _sandbox(
        tmp_path,
        MASTER_TITLES[:5],
        [_themed(1) for _ in MASTER_TITLES[:5]],
    )
    dest = filler.gather_and_fill(sandbox)
    texts = _job_paragraphs(dest)
    joined = "\n".join(texts)
    assert "{{" not in joined
    assert "Projects" in texts
    assert "Career Forge" in texts
    assert "Automated job-search pipeline." in texts
    assert "Hidden Project" not in joined
    edu = texts.index("Education")
    proj = texts.index("Projects")
    skills = texts.index("Skills")
    assert edu < proj < skills
    markers = _right_column_markers(dest)
    assert markers.count("TOPBORDER") == 2
    assert (
        markers.index("Education")
        < markers.index("TOPBORDER")
        < markers.index("Projects")
        < markers.index("Career Forge")
        < markers.index("Automated job-search pipeline.")
        < markers.index("TOPBORDER", markers.index("Projects"))
        < markers.index("Skills")
    )
    title_found = False
    for para in filler.iter_paragraphs(Document(str(dest))):
        if para.text.strip() != "Career Forge":
            continue
        title_found = True
        assert any(run.bold for run in para.runs if run.text)
    assert title_found


def test_populated_project_with_optional_url_renders_url(
    tmp_path: Path, monkeypatch
) -> None:
    proj_path = tmp_path / "locked-projects.json"
    proj_path.write_text(
        json.dumps(
            {
                "id": "prj_001",
                "1": {
                    "title": "Career Forge",
                    "url": "https://github.com/example/job-forge",
                    "summary": "Automated job-search pipeline.",
                },
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(filler, "PROJECTS_PATH", proj_path)
    sandbox = _sandbox(
        tmp_path,
        MASTER_TITLES[:5],
        [_themed(1) for _ in MASTER_TITLES[:5]],
    )
    dest = filler.gather_and_fill(sandbox)
    texts = _job_paragraphs(dest)
    joined = "\n".join(texts)
    assert "{{" not in joined
    assert "Projects" in texts
    assert "Career Forge" in texts
    assert "https://github.com/example/job-forge" in texts
    assert "Automated job-search pipeline." in texts
    markers = _right_column_markers(dest)
    assert markers.count("TOPBORDER") == 2
    assert (
        markers.index("Education")
        < markers.index("TOPBORDER")
        < markers.index("Projects")
        < markers.index("Career Forge")
        < markers.index("https://github.com/example/job-forge")
        < markers.index("Automated job-search pipeline.")
        < markers.index("TOPBORDER", markers.index("Projects"))
        < markers.index("Skills")
    )


def test_n10_success_keeps_all_printed_roles(tmp_path: Path) -> None:
    sandbox = _sandbox(tmp_path, MASTER_TITLES, [_themed(1) for _ in MASTER_TITLES])
    dest = filler.gather_and_fill(sandbox)
    texts = _job_paragraphs(dest)
    joined = "\n".join(texts)
    assert "Application Development Systems Analyst" in joined
    assert "Associate Technical Support Analyst & Account Coordinator" in joined
    assert "IT Systems Consultant" not in joined
    assert "Director of Information Technology" not in joined
    assert "{{" not in joined


def test_theme_bold_uses_first_colon_only(tmp_path: Path) -> None:
    sandbox = _sandbox(
        tmp_path,
        MASTER_TITLES[:5],
        [
            ["Architecture Governance: Established an ARB: still body."],
            _themed(1),
            _themed(1),
            _themed(1),
            _themed(1),
        ],
    )
    dest = filler.gather_and_fill(sandbox)
    found = False
    for para in filler.iter_paragraphs(Document(str(dest))):
        if "Established an ARB" not in para.text:
            continue
        found = True
        runs = [run for run in para.runs if run.text]
        assert runs[0].text == "Architecture Governance:"
        assert runs[0].bold is True
        body = "".join(run.text for run in runs[1:])
        assert body == " Established an ARB: still body."
        assert all(run.bold in (False, None) for run in runs[1:])
        b_vals = []
        for run in runs[1:]:
            r_pr = run._r.find(qn("w:rPr"))
            if r_pr is None:
                continue
            for child in r_pr:
                if child.tag == qn("w:b"):
                    b_vals.append(child.get(qn("w:val")))
        assert all(val in {None, "0", "false"} for val in b_vals)
    assert found
