"""Locked and re-derived templates must expose JOB_HISTORY slots 1-10."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document

_DOCX_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_DOCX_DIR))

import derive_template
import fill_resume as filler


def test_locked_template_has_complete_slot_inventory() -> None:
    tokens = filler.collect_template_tokens(Document(str(derive_template.OUTPUT)))
    missing = derive_template.job_slot_token_inventory() - tokens
    assert missing == set()


def test_locked_template_has_contact_slot_inventory() -> None:
    tokens = filler.collect_template_tokens(Document(str(derive_template.OUTPUT)))
    missing = derive_template.contact_slot_token_inventory() - tokens
    assert missing == set()
    assert "CONTACT_LINKS_1_URL" in tokens


def test_locked_template_has_project_1_between_education_and_skills() -> None:
    doc = Document(str(derive_template.OUTPUT))
    tokens = filler.collect_template_tokens(doc)
    missing = derive_template.project_slot_token_inventory() - tokens
    assert missing == set()
    extra = {
        token
        for token in tokens
        if token.startswith("PROJECTS_") and token not in derive_template.project_slot_token_inventory()
    }
    assert extra == set()
    assert derive_template.right_column_section_markers(doc) == [
        "Education",
        "TOPBORDER",
        "Projects",
        "PROJECTS_1_TITLE",
        "PROJECTS_1_URL",
        "PROJECTS_1_SUMMARY",
        "TOPBORDER",
        "Skills",
    ]


@pytest.mark.skipif(
    not derive_template.SAMPLE.is_file(),
    reason="layout sample is not shipped in the public repository",
)
def test_rederive_emits_complete_slot_inventory(tmp_path: Path, monkeypatch) -> None:
    dest = tmp_path / "locked-resume-template.docx"
    original = derive_template.OUTPUT
    monkeypatch.setattr(derive_template, "OUTPUT", dest)
    derive_template.derive()
    tokens = filler.collect_template_tokens(Document(str(dest)))
    missing = derive_template.job_slot_token_inventory() - tokens
    assert missing == set()
    missing_contact = derive_template.contact_slot_token_inventory() - tokens
    assert missing_contact == set()
    assert "CONTACT_LINKS_1_URL" in tokens
    assert dest.is_file()
    assert original != dest


@pytest.mark.skipif(
    not derive_template.SAMPLE.is_file(),
    reason="layout sample is not shipped in the public repository",
)
def test_rederive_emits_project_1_between_education_and_skills(
    tmp_path: Path, monkeypatch
) -> None:
    dest = tmp_path / "locked-resume-template.docx"
    monkeypatch.setattr(derive_template, "OUTPUT", dest)
    derive_template.derive()
    doc = Document(str(dest))
    tokens = filler.collect_template_tokens(doc)
    missing = derive_template.project_slot_token_inventory() - tokens
    assert missing == set()
    extra = {
        token
        for token in tokens
        if token.startswith("PROJECTS_") and token not in derive_template.project_slot_token_inventory()
    }
    assert extra == set()
    assert derive_template.right_column_section_markers(doc) == [
        "Education",
        "TOPBORDER",
        "Projects",
        "PROJECTS_1_TITLE",
        "PROJECTS_1_URL",
        "PROJECTS_1_SUMMARY",
        "TOPBORDER",
        "Skills",
    ]
