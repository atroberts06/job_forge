#!/usr/bin/env python3
"""Fill locked-cover-letter-template.docx from sandbox markdown (ADR-053).

Discovers {{TOKENS}} in the template, binds locked contact fields, and overlays
section blocks parsed from tailored-cover-letter.md. Never writes .ai/guardrails/.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from html import unescape
from pathlib import Path

from lxml import etree

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_MODEL = REPO_ROOT / ".ai/history/architectural-decisions/cover-letter-data-model.json"
CONTACT_PATH = REPO_ROOT / ".ai/guardrails/locked-contact.json"
SKILLS_PATH = REPO_ROOT / ".ai/guardrails/locked-skills.json"
JOB_HISTORY_PATH = REPO_ROOT / ".ai/guardrails/locked-job-history.json"
SUMMARY_PATH = REPO_ROOT / ".ai/guardrails/locked-professional-summary.json"
TEMPLATE_PATH = REPO_ROOT / ".ai/guardrails/locked-cover-letter-template.docx"

REQUIRED_LOCKED = {
    "contact": CONTACT_PATH,
    "skills": SKILLS_PATH,
    "job_history": JOB_HISTORY_PATH,
    "professional_summary": SUMMARY_PATH,
}

TOKEN_RE = re.compile(r"\{\{([^{}]+)\}\}")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
REQUIRED_SECTIONS = (
    "Opener",
    "Body 1",
    "Body 2",
    "Call to Action",
    "Closing",
)
METADATA_KEYS = ("Date", "Recipient", "Subject", "Greeting")
SECTION_TOKEN_MAP = {
    "Opener": "COVER_LETTER_OPENER",
    "Body 1": "COVER_LETTER_BODY_1",
    "Body 2": "COVER_LETTER_BODY_2",
    "Call to Action": "COVER_LETTER_CALL_TO_ACTION",
}
CONTACT_REQUIRED = (
    "first_name",
    "last_name",
    "email",
    "phone",
    "city",
    "state",
    "postal_code",
)
_DOCX_XML_PREFIXES = ("word/", "customxml/")


class FillError(Exception):
    """Human-readable fill abort."""


def load_json(path: Path) -> dict:
    if not path.is_file():
        raise FillError(f"missing required file: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise FillError(f"{path}: invalid JSON — {exc}") from exc


def require_fields(obj: dict, fields: list[str] | tuple[str, ...], label: str) -> None:
    missing = [field for field in fields if field not in obj or obj[field] in (None, "")]
    if missing:
        raise FillError(f"{label} missing required field(s): {', '.join(missing)}")


def utc_today_long() -> str:
    now = datetime.now(timezone.utc)
    return f"{now.strftime('%B')} {now.day}, {now.year}"


def company_title_case(company_slug: str) -> str:
    return "_".join(part.capitalize() for part in company_slug.split("-"))


def output_name(sandbox: Path) -> str:
    company_slug = sandbox.parent.name
    return f"John_Doe_Cover_Letter_{company_title_case(company_slug)}.docx"


def iter_paragraphs(parent):
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    elif isinstance(parent, Table):
        parent_elm = parent._tbl
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            seen: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    tc_id = id(cell._tc)
                    if tc_id in seen:
                        continue
                    seen.add(tc_id)
                    yield from iter_paragraphs(cell)


def collect_template_tokens(doc) -> set[str]:
    tokens: set[str] = set()
    for para in iter_paragraphs(doc):
        tokens.update(TOKEN_RE.findall(para.text))
    return tokens


def _preserve_w_t_spaces(run) -> None:
    for node in run._r.iter(qn("w:t")):
        if node.text and "  " in node.text:
            node.set(XML_SPACE, "preserve")


def _copy_rpr(r_pr):
    return deepcopy(r_pr) if r_pr is not None else None


def _run_r_pr(run):
    r_pr = run._r.find(qn("w:rPr"))
    return _copy_rpr(r_pr)


def _rpr_key(r_pr) -> str:
    if r_pr is None:
        return ""
    xml = getattr(r_pr, "xml", None)
    if xml is not None:
        return str(xml)
    return etree.tostring(r_pr, encoding="unicode")


def _text_runs(paragraph: Paragraph) -> list[Run]:
    p_el = paragraph._p
    runs: list[Run] = []
    for r_el in p_el.iter(qn("w:r")):
        parent_p = r_el
        while parent_p is not None and parent_p.tag != qn("w:p"):
            parent_p = parent_p.getparent()
        if parent_p is p_el:
            runs.append(Run(r_el, paragraph))
    return runs


def paragraph_run_parts(paragraph: Paragraph) -> list[tuple[str, object]]:
    return [(run.text or "", _run_r_pr(run)) for run in _text_runs(paragraph)]


def _majority_part_index(
    owners: list[int], start: int, end: int, parts: list[tuple[str, object]]
) -> int:
    rpr_count: dict[str, int] = {}
    rpr_first_idx: dict[str, int] = {}
    for idx in owners[start:end]:
        key = _rpr_key(parts[idx][1])
        rpr_count[key] = rpr_count.get(key, 0) + 1
        rpr_first_idx.setdefault(key, idx)
    best_key = max(rpr_count, key=lambda key: (rpr_count[key], -rpr_first_idx[key]))
    return rpr_first_idx[best_key]


def replace_tokens_in_parts(
    parts: list[tuple[str, object]], mapping: dict[str, str]
) -> list[tuple[str, object]]:
    if not parts:
        return parts
    full = "".join(text for text, _ in parts)
    if "{{" not in full:
        return parts
    owners: list[int] = []
    for index, (text, _) in enumerate(parts):
        owners.extend([index] * len(text))
    spans: list[tuple[int, int, str]] = []
    for match in TOKEN_RE.finditer(full):
        key = "{{" + match.group(1) + "}}"
        if key in mapping:
            spans.append((match.start(), match.end(), mapping[key]))
    if not spans:
        return parts
    span_at = {start: (end, value) for start, end, value in spans}
    buckets: list[list[str]] = [[] for _ in parts]
    index = 0
    length = len(full)
    while index < length:
        if index in span_at:
            end, value = span_at[index]
            winner = _majority_part_index(owners, index, end, parts)
            buckets[winner].append(value)
            index = end
            continue
        buckets[owners[index]].append(full[index])
        index += 1
    return [("".join(buckets[i]), parts[i][1]) for i in range(len(parts))]


def apply_parts_to_paragraph(paragraph: Paragraph, parts: list[tuple[str, object]]) -> None:
    runs = _text_runs(paragraph)
    for run, (text, _) in zip(runs, parts):
        run.text = text
        _preserve_w_t_spaces(run)
    if len(runs) < len(parts):
        raise FillError("placeholder run rewrite lost template runs")


def replace_placeholders(paragraph: Paragraph, mapping: dict[str, str]) -> None:
    text = paragraph.text
    if "{{" not in text:
        return
    parts = replace_tokens_in_parts(paragraph_run_parts(paragraph), mapping)
    filled = "".join(chunk for chunk, _ in parts)
    if filled != text:
        apply_parts_to_paragraph(paragraph, parts)


def parse_job_heading(path: Path) -> tuple[str, str]:
    if not path.is_file():
        raise FillError(f"missing required file: {path}")
    for raw in path.read_text(encoding="utf-8").splitlines():
        stripped = raw.strip()
        if stripped.startswith("# "):
            heading = stripped[2:].strip()
            if " - " in heading:
                company, title = heading.split(" - ", 1)
                return company.strip(), title.strip()
            return heading, ""
    raise FillError(f"{path}: missing H1 company - title heading")


def parse_job_id(sandbox: Path) -> str:
    marker = sandbox / "job-id.txt"
    if not marker.is_file():
        return ""
    text = marker.read_text(encoding="utf-8").strip()
    if not text:
        return ""
    return text.rsplit(":", 1)[-1].strip()


def parse_cover_letter(path: Path) -> dict[str, str]:
    if not path.is_file():
        raise FillError(f"missing required file: {path}")
    metadata: dict[str, str] = {}
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        stripped = raw.strip()
        if stripped.startswith("## "):
            current = stripped[3:].strip()
            sections.setdefault(current, [])
            continue
        meta = re.match(r"^- \*\*(.+?):\*\*\s*(.*)$", stripped)
        if meta and current is None:
            metadata[meta.group(1).strip()] = meta.group(2).strip()
            continue
        if current is None:
            continue
        if stripped:
            sections[current].append(stripped)
    missing = [name for name in REQUIRED_SECTIONS if name not in sections or not sections[name]]
    if missing:
        raise FillError(
            f"{path}: missing required section(s): {', '.join(missing)}"
        )
    parsed = {key: metadata.get(key, "").strip() for key in METADATA_KEYS}
    parsed["Opener"] = " ".join(sections["Opener"])
    parsed["Body 1"] = " ".join(sections["Body 1"])
    parsed["Body 2"] = " ".join(sections["Body 2"])
    parsed["Call to Action"] = " ".join(sections["Call to Action"])
    closing_lines = sections["Closing"]
    parsed["Sign-off"] = closing_lines[0].rstrip()
    leftover_parts = [
        parsed[name]
        for name in (*METADATA_KEYS, "Opener", "Body 1", "Body 2", "Call to Action", "Sign-off")
        if "{{" in parsed[name]
    ]
    if leftover_parts:
        raise FillError(
            f"{path}: unresolved nested token(s) remain in cover letter markdown"
        )
    return parsed


def contact_mapping(contact: dict) -> dict[str, str]:
    require_fields(contact, CONTACT_REQUIRED, "contact")
    links = contact.get("links") or []
    url = ""
    if isinstance(links, list) and links:
        first = links[0]
        if isinstance(first, dict):
            url = str(first.get("url") or "").strip()
    if not url:
        raise FillError("contact missing required field(s): links[0].url")
    first = str(contact["first_name"]).strip()
    last = str(contact["last_name"]).strip()
    return {
        "{{CONTACT_FIRST_NAME}}": first,
        "{{CONTACT_LAST_NAME}}": last,
        "{{CONTACT_FULL_NAME}}": f"{first} {last}".strip(),
        "{{CONTACT_EMAIL}}": str(contact["email"]).strip(),
        "{{CONTACT_PHONE}}": str(contact["phone"]).strip(),
        "{{CONTACT_CITY}}": str(contact["city"]).strip(),
        "{{CONTACT_STATE}}": str(contact["state"]).strip(),
        "{{CONTACT_POSTAL_CODE}}": str(contact["postal_code"]).strip(),
        "{{CONTACT_LINKS_1_URL}}": url,
    }


def build_mapping(
    tokens: set[str],
    *,
    contact: dict,
    cover: dict[str, str],
    company: str,
    title: str,
    job_id: str,
) -> dict[str, str]:
    date_value = cover.get("Date") or utc_today_long()
    recipient = cover.get("Recipient") or "Hiring Team"
    subject = cover.get("Subject") or f"RE: {title}, Req #{job_id}, {date_value}"
    greeting = cover.get("Greeting") or f"Dear {recipient},"
    mapping = contact_mapping(contact)
    mapping.update(
        {
            "{{CURRENT_DATE}}": date_value,
            "{{COMPANY_NAME}}": company,
            "{{JOB_TITLE}}": title,
            "{{JOB_ID}}": job_id,
            "{{REF_NUM}}": job_id,
            "{{RECIPIENT_NAME}}": recipient,
            "{{COVER_LETTER_SUBJECT}}": subject,
            "{{COVER_LETTER_GREETING}}": greeting,
            "{{COVER_LETTER_OPENER}}": cover["Opener"],
            "{{COVER_LETTER_BODY_1}}": cover["Body 1"],
            "{{COVER_LETTER_BODY_2}}": cover["Body 2"],
            "{{COVER_LETTER_CALL_TO_ACTION}}": cover["Call to Action"],
            "{{COVER_LETTER_SIGN_OFF}}": cover["Sign-off"],
        }
    )
    unknown = []
    for body in sorted(tokens):
        key = "{{" + body + "}}"
        if key not in mapping:
            unknown.append(key)
    if unknown:
        raise FillError(f"unknown placeholder: {', '.join(unknown)}")
    return {key: mapping[key] for key in ( "{{" + body + "}}" for body in tokens)}


def unresolved_docx_placeholders(docx_path: Path) -> list[str]:
    hits: list[str] = []
    with zipfile.ZipFile(docx_path) as archive:
        for name in archive.namelist():
            lower = name.replace("\\", "/").lower()
            if not lower.endswith(".xml"):
                continue
            if not lower.startswith(_DOCX_XML_PREFIXES):
                continue
            raw = archive.read(name).decode("utf-8")
            visible = unescape(re.sub(r"<[^>]+>", "", raw))
            if "{{" in visible:
                hits.append(name)
    return hits


def fill_document(template: Path, mapping: dict[str, str], dest: Path) -> None:
    shutil.copy2(template, dest)
    doc = Document(str(dest))
    for para in iter_paragraphs(doc):
        if "{{" not in para.text:
            continue
        replace_placeholders(para, mapping)
        after = para.text
        if "{{" in after and "}}" in after:
            raise FillError(f"unresolved placeholder remaining: {after}")
    doc.save(str(dest))
    leftover = unresolved_docx_placeholders(dest)
    if leftover:
        raise FillError(
            f"leftover template placeholder '{{{{' in {', '.join(leftover)}"
        )


def gather_and_fill(sandbox: Path) -> Path:
    load_json(DATA_MODEL)
    contact: dict | None = None
    for label, path in REQUIRED_LOCKED.items():
        payload = load_json(path)
        if label == "contact":
            contact = payload
    if contact is None:
        raise FillError(f"missing required file: {CONTACT_PATH}")
    company, title = parse_job_heading(sandbox / "job-description.md")
    cover = parse_cover_letter(sandbox / "tailored-cover-letter.md")
    job_id = parse_job_id(sandbox)
    if not TEMPLATE_PATH.is_file():
        raise FillError(f"missing required file: {TEMPLATE_PATH}")
    template_doc = Document(str(TEMPLATE_PATH))
    tokens = collect_template_tokens(template_doc)
    mapping = build_mapping(
        tokens,
        contact=contact,
        cover=cover,
        company=company,
        title=title,
        job_id=job_id,
    )
    dest = sandbox / output_name(sandbox)
    try:
        fill_document(TEMPLATE_PATH, mapping, dest)
    except PermissionError as exc:
        raise FillError(
            f"cannot write {dest} - close the file if it is open in Word"
        ) from exc
    return dest


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Fill locked-cover-letter-template.docx for a resume-tailor sandbox."
    )
    parser.add_argument(
        "--sandbox",
        required=True,
        help="Path to .ai/history/resume-tailor/{company}/{position}",
    )
    args = parser.parse_args()
    sandbox = Path(args.sandbox)
    if not sandbox.is_dir():
        print(f"ERROR: sandbox directory not found: {sandbox}", file=sys.stderr)
        return 1
    try:
        dest = gather_and_fill(sandbox)
    except FillError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(f"wrote {dest}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
