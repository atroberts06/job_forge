"""Locked cover letter template in .ai/guardrails exposes the 8-theme slot inventory."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

_DOCX_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_DOCX_DIR))

import derive_cover_letter_template as derive
import fill_cover_letter as filler


def _nonempty_paragraphs(path: Path) -> list[str]:
    return [
        para.text.strip()
        for para in Document(str(path)).paragraphs
        if para.text.strip()
    ]


def test_locked_template_has_section_and_header_inventory() -> None:
    tokens = filler.collect_template_tokens(Document(str(derive.OUTPUT)))
    missing = derive.expected_template_tokens() - tokens
    assert missing == set()
    extra_sections = derive.section_slot_token_inventory() - tokens
    assert extra_sections == set()
    texts = _nonempty_paragraphs(derive.OUTPUT)
    assert texts.count("{{CONTACT_FULL_NAME}}") == 2
    assert "{{COVER_LETTER_CALL_TO_ACTION}}" in texts


def test_locked_template_matches_golden_fixture() -> None:
    assert derive.OUTPUT.is_file()
    assert derive.GOLDEN.is_file()
    locked_tokens = filler.collect_template_tokens(Document(str(derive.OUTPUT)))
    golden_tokens = filler.collect_template_tokens(Document(str(derive.GOLDEN)))
    assert locked_tokens == golden_tokens
    assert _nonempty_paragraphs(derive.OUTPUT) == _nonempty_paragraphs(derive.GOLDEN)
