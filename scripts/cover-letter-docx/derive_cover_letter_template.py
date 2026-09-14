"""One-time derivation of locked-cover-letter-template.docx from the POC layout.

Layout only: replace identity/content with section-block placeholders and strip
vendor custom properties. Does not invent jobs, skills, titles, dates, or metrics.
"""

from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

REPO_ROOT = Path(__file__).resolve().parents[2]
SAMPLE = REPO_ROOT / "scripts/cover-letter-docx/tests/fixtures/layout-sample.docx"
OUTPUT = REPO_ROOT / ".ai/guardrails/locked-cover-letter-template.docx"
GOLDEN = (
    REPO_ROOT
    / ".github/scripts/resume-tailor/tests/fixtures/artifacts"
    / "locked-cover-letter-template.docx"
)

SECTION_TOKENS = {
    "COVER_LETTER_SUBJECT",
    "COVER_LETTER_GREETING",
    "COVER_LETTER_OPENER",
    "COVER_LETTER_BODY_1",
    "COVER_LETTER_BODY_2",
    "COVER_LETTER_CALL_TO_ACTION",
    "COVER_LETTER_SIGN_OFF",
}

CONTACT_TOKENS = {
    "CONTACT_FIRST_NAME",
    "CONTACT_LAST_NAME",
    "CONTACT_FULL_NAME",
    "CONTACT_EMAIL",
    "CONTACT_PHONE",
    "CONTACT_CITY",
    "CONTACT_STATE",
    "CONTACT_POSTAL_CODE",
    "CONTACT_LINKS_1_URL",
}

RUNTIME_TOKENS = {
    "CURRENT_DATE",
    "COMPANY_NAME",
    "JOB_TITLE",
    "JOB_ID",
    "REF_NUM",
    "RECIPIENT_NAME",
}

OPENER_MARKER = "As an accomplished and seasoned professional"
BODY_1_MARKER = "I offer a combination of unique skills"
BODY_2_MARKER = "As my attached resume indicates"
CTA_MARKER = "I am eager to discuss the possibility"


def iter_paragraphs(parent):
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    elif isinstance(parent, Table):
        parent_elm = parent._tbl
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            seen: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    tc_id = id(cell._tc)
                    if tc_id in seen:
                        continue
                    seen.add(tc_id)
                    yield from iter_paragraphs(cell)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.add_run(text)


def clone_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def section_slot_token_inventory() -> set[str]:
    return set(SECTION_TOKENS)


def contact_slot_token_inventory() -> set[str]:
    return set(CONTACT_TOKENS)


def runtime_slot_token_inventory() -> set[str]:
    return set(RUNTIME_TOKENS)


def expected_template_tokens() -> set[str]:
    """Tokens the derived Word template must expose (section + header + date)."""
    return SECTION_TOKENS | {
        "CONTACT_FULL_NAME",
        "CURRENT_DATE",
    }


def strip_vendor_metadata(docx_path: Path) -> None:
    buffer = BytesIO(docx_path.read_bytes())
    with zipfile.ZipFile(buffer, "r") as zin:
        payload = {name: zin.read(name) for name in zin.namelist()}

    payload.pop("docProps/custom.xml", None)

    rels_path = "_rels/.rels"
    if rels_path in payload:
        ET = __import__("xml.etree.ElementTree", fromlist=["ET"])
        root = ET.fromstring(payload[rels_path])
        ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        changed = False
        for rel in list(root.findall("r:Relationship", ns)):
            target = rel.get("Target", "")
            rel_type = rel.get("Type", "")
            if "custom.xml" in target or "custom-properties" in rel_type:
                root.remove(rel)
                changed = True
        if changed:
            payload[rels_path] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True
            )

    ctypes_path = "[Content_Types].xml"
    if ctypes_path in payload:
        ET = __import__("xml.etree.ElementTree", fromlist=["ET"])
        root = ET.fromstring(payload[ctypes_path])
        ns = {"t": "http://schemas.openxmlformats.org/package/2006/content-types"}
        changed = False
        for node in list(root.findall("t:Override", ns)):
            if node.get("PartName") == "/docProps/custom.xml":
                root.remove(node)
                changed = True
        if changed:
            payload[ctypes_path] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True
            )

    if "docProps/core.xml" in payload:
        ET = __import__("xml.etree.ElementTree", fromlist=["ET"])
        root = ET.fromstring(payload["docProps/core.xml"])
        nsmap = {
            "dc": "http://purl.org/dc/elements/1.1/",
            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        }
        for tag in ("dc:title", "dc:creator", "cp:lastModifiedBy"):
            prefix, local = tag.split(":")
            for node in list(root.findall(f"{prefix}:{local}", nsmap)):
                node.text = ""
        payload["docProps/core.xml"] = ET.tostring(
            root, encoding="utf-8", xml_declaration=True
        )

    with zipfile.ZipFile(docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in payload.items():
            zout.writestr(name, data)


def _inject_contact_header(name_para: Paragraph) -> None:
    set_paragraph_text(name_para, "{{CONTACT_FULL_NAME}}")


def derive() -> None:
    if not SAMPLE.is_file():
        raise SystemExit(f"missing layout sample: {SAMPLE}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SAMPLE, OUTPUT)

    doc = Document(str(OUTPUT))
    name_para: Paragraph | None = None
    body2_para: Paragraph | None = None
    sign_para: Paragraph | None = None

    for para in iter_paragraphs(doc):
        text = para.text.strip()
        collapsed = text.replace(" ", "")
        if text == "John Doe" and name_para is None:
            name_para = para
            continue
        if text.startswith("June 30, 2026") or text == "June 30, 2026":
            set_paragraph_text(para, "{{CURRENT_DATE}}")
            continue
        if text.startswith("RE:"):
            set_paragraph_text(para, "{{COVER_LETTER_SUBJECT}}")
            continue
        if text.startswith("Dear "):
            set_paragraph_text(para, "{{COVER_LETTER_GREETING}}")
            continue
        if OPENER_MARKER in text:
            set_paragraph_text(para, "{{COVER_LETTER_OPENER}}")
            continue
        if BODY_1_MARKER in text:
            set_paragraph_text(para, "{{COVER_LETTER_BODY_1}}")
            continue
        if BODY_2_MARKER in text:
            set_paragraph_text(para, "{{COVER_LETTER_BODY_2}}")
            body2_para = para
            if CTA_MARKER in text:
                cta = clone_paragraph_after(para)
                set_paragraph_text(cta, "{{COVER_LETTER_CALL_TO_ACTION}}")
            continue
        if text.startswith("I am eager to discuss") and body2_para is not None:
            set_paragraph_text(para, "{{COVER_LETTER_CALL_TO_ACTION}}")
            continue
        if "Sincerely" in text:
            set_paragraph_text(para, "{{COVER_LETTER_SIGN_OFF}}")
            sign_para = para
            if "John Doe" in text or "JohnDoe" in collapsed:
                name_close = clone_paragraph_after(para)
                set_paragraph_text(name_close, "{{CONTACT_FULL_NAME}}")
            continue
        if text == "John Doe" and sign_para is not None and para is not name_para:
            set_paragraph_text(para, "{{CONTACT_FULL_NAME}}")

    if name_para is None:
        raise SystemExit("cannot derive cover letter: name header paragraph missing")
    _inject_contact_header(name_para)

    if body2_para is not None:
        joined = body2_para.text
        if CTA_MARKER in joined and "{{COVER_LETTER_CALL_TO_ACTION}}" not in "".join(
            p.text for p in iter_paragraphs(doc)
        ):
            cta = clone_paragraph_after(body2_para)
            set_paragraph_text(cta, "{{COVER_LETTER_CALL_TO_ACTION}}")

    texts = [p.text.strip() for p in iter_paragraphs(doc) if p.text.strip()]
    required = [
        "{{CONTACT_FULL_NAME}}",
        "{{CURRENT_DATE}}",
        "{{COVER_LETTER_SUBJECT}}",
        "{{COVER_LETTER_GREETING}}",
        "{{COVER_LETTER_OPENER}}",
        "{{COVER_LETTER_BODY_1}}",
        "{{COVER_LETTER_BODY_2}}",
        "{{COVER_LETTER_CALL_TO_ACTION}}",
        "{{COVER_LETTER_SIGN_OFF}}",
    ]
    missing = [token for token in required if not any(token in text for text in texts)]
    if missing:
        raise SystemExit(f"derive missed tokens: {missing}; paragraphs={texts}")

    doc.save(str(OUTPUT))
    strip_vendor_metadata(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    derive()
